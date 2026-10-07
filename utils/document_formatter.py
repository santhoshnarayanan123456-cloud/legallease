from __future__ import annotations

import html
import re

from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.shared import Inches, Pt

from fpdf import FPDF


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",
        "\u2019": "'",

        "\u201c": '"',
        "\u201d": '"',

        "\u2013": "-",
        "\u2014": "-",

        "\u00a0": " ",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    text = re.sub(
        r"[ \t]+\n",
        "\n",
        text
    )

    return text.strip()


def is_heading(line: str) -> bool:

    stripped = line.strip()

    if not stripped:

        return False

    if re.match(
        r"^\d+[\.\)]\s+",
        stripped
    ):

        return True

    if (
        stripped.isupper()
        and len(stripped) < 100
    ):

        return True

    heading_names = {
        "drafting note",
        "signatures",
        "signature",
        "terms",
    }

    if stripped.lower() in heading_names:

        return True

    return False


def format_docx(
    text: str,
    doc_type: str,
    logo_path: Optional[str] = None
) -> bytes:

    text = sanitize_text(text)

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # Logo

    if (
        logo_path
        and Path(logo_path).exists()
    ):

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            logo_path,
            width=Inches(1.35)
        )

    # Title

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    # Document body

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:

            document.add_paragraph()

            continue

        if is_heading(line):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

            continue

        if line.startswith(
            ("- ", "* ")
        ):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:].strip()
            )

            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)

        paragraph.add_run(line)

    # Footer

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review before signing."
    )

    footer_run.italic = True

    # Convert to bytes

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type: str,
        logo_path: Optional[str] = None
    ):

        super().__init__()

        self.doc_type = doc_type
        self.logo_path = logo_path

    def header(self):

        if (
            self.logo_path
            and Path(self.logo_path).exists()
        ):

            self.image(
                self.logo_path,
                x=90,
                y=8,
                w=30
            )

            self.set_y(40)

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            8,
            self.doc_type,
            align="C"
        )

        self.ln(10)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft. "
            "Review before signing.",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: Optional[str] = None
) -> bytes:

    text = sanitize_text(text)

    pdf = LegalEasePDF(
        doc_type,
        logo_path
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        size=11
    )

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:

            pdf.ln(4)

            continue

        if is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                line,
                 new_x="LMARGIN",
                 new_y="NEXT",
            )

            pdf.set_font(
                "Helvetica",
                size=11
            )

        elif line.startswith(
            ("- ", "* ")
        ):

            pdf.multi_cell(
                0,
                6,
                "  " + line
            )

        else:

            pdf.multi_cell(
                0,
                6,
                line
            )

    return bytes(
        pdf.output()
    )


def format_html_preview(
    text: str
) -> str:

    safe_text = html.escape(
        sanitize_text(text)
    )

    lines = safe_text.splitlines()

    rendered = []

    for line in lines:

        if not line.strip():

            rendered.append(
                "<div style='height:8px'></div>"
            )

        elif is_heading(
            html.unescape(line)
        ):

            rendered.append(
                f"<h4>{line}</h4>"
            )

        elif line.startswith(
            ("- ", "* ")
        ):

            rendered.append(
                f"<div>&bull; {line[2:]}</div>"
            )

        else:

            rendered.append(
                f"<p>{line}</p>"
            )

    return f"""
    <div style="
        background:#111827;
        color:#f3f4f6;
        padding:24px;
        border-radius:12px;
        max-height:620px;
        overflow-y:auto;
        line-height:1.65;
        font-family:Georgia, 'Times New Roman', serif;
    ">
        {''.join(rendered)}
    </div>
    """